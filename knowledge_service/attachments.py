"""Office extraction and page-level OCR for official attachments."""
from io import BytesIO
from pathlib import Path
import tempfile
import subprocess
import os
import threading

_ocr_engine=None
_ocr_lock=threading.Lock()


def recognize(image):
    global _ocr_engine
    from rapidocr_onnxruntime import RapidOCR
    with _ocr_lock:
        if _ocr_engine is None:
            _ocr_engine=RapidOCR(intra_op_num_threads=4, inter_op_num_threads=1)
        result,_=_ocr_engine(image)
    return '\n'.join(line[1] for line in result or [])


def ocr_images(urls):
    import requests
    pages=[]
    for i,url in enumerate(urls):
        response=requests.get(url,timeout=(10,30));response.raise_for_status()
        pages.append((i+1,recognize(response.content)))
    return pages


def office_text(data, suffix):
    if suffix=='.docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document=Document(BytesIO(data));parts=[]
        for node in document.element.body:
            if node.tag==qn('w:p'):
                parts.append(Paragraph(node,document).text)
            elif node.tag==qn('w:tbl'):
                parts.extend(' | '.join(cell.text for cell in row.cells) for row in Table(node,document).rows)
        return '\n'.join(parts)
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        workbook=load_workbook(BytesIO(data),read_only=True,data_only=True)
        parts=[]
        for sheet in workbook:
            parts.append('工作表：'+sheet.title)
            parts.extend(' | '.join('' if v is None else str(v) for v in row) for row in sheet.iter_rows(values_only=True))
        workbook.close()
        return '\n'.join(parts)
    if suffix in {'.doc','.xls'}:
        executable=os.environ['KNOWLEDGE_LIBREOFFICE']
        target_suffix='.docx' if suffix=='.doc' else '.xlsx'
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);source=root/('source'+suffix);source.write_bytes(data)
            subprocess.run([executable,'--headless','-env:UserInstallation='+ (root/'profile').as_uri(),
                '--convert-to',target_suffix[1:],'--outdir',directory,str(source)],check=True,timeout=90,
                stdout=subprocess.PIPE,stderr=subprocess.PIPE)
            return office_text((root/('source'+target_suffix)).read_bytes(),target_suffix)
    raise ValueError('Unsupported attachment format')


def ocr_pdf(data, page_indexes):
    import pypdfium2 as pdfium
    import numpy as np
    output={}
    document=pdfium.PdfDocument(data)
    for index in page_indexes:
        page=document[index]
        bitmap=page.render(scale=2)
        output[index]=recognize(np.asarray(bitmap.to_pil()))
        bitmap.close();page.close()
    document.close()
    return output
